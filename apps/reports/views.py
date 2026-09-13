import csv
from io import BytesIO
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.utils import timezone

from apps.ncr.models import NCR
from apps.capa.models import CAPA
from apps.actions.models import Action
from apps.risks.models import Risk
from apps.inspections.models import Inspection
from apps.calibration.models import Equipment
from apps.training.models import TrainingRecord
from apps.audits.models import AuditFinding
from apps.departments.models import Department

@login_required
def report_index(request):
    departments = Department.objects.all()
    return render(request, 'reports/index.html', {
        'departments': departments,
        'today': timezone.now().date(),
    })

@login_required
def export_dataset(request):
    module = request.GET.get('module', 'ncr')
    export_format = request.GET.get('format', 'xlsx')
    dept_id = request.GET.get('department')
    start_date = request.GET.get('start_date')
    end_date = request.GET.get('end_date')

    # Prepare datasets based on module
    headers, rows, filename_prefix = get_module_data(module, dept_id, start_date, end_date)
    timestamp = timezone.now().strftime('%Y%m%d_%H%M%S')
    filename = f"{filename_prefix}_{timestamp}.{export_format}"

    if export_format == 'csv':
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        writer = csv.writer(response)
        writer.writerow(headers)
        for row in rows:
            writer.writerow(row)
        return response

    elif export_format == 'xlsx':
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = module.upper()

        # Styles
        header_fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Calibri", size=10)
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        # Write Header
        ws.append(headers)
        for col_num, cell in enumerate(ws[1], 1):
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Write Data
        for row in rows:
            ws.append(row)

        # Apply borders and auto-fit columns
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
            for cell in row:
                cell.font = regular_font
                cell.border = thin_border

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response

    return HttpResponse("Invalid format", status=400)


def get_module_data(module, dept_id=None, start_date=None, end_date=None):
    if module == 'ncr':
        qs = NCR.objects.select_related('department', 'reported_by', 'process').all()
        if dept_id:
            qs = qs.filter(department_id=dept_id)
        if start_date:
            qs = qs.filter(created_at__date__gte=start_date)
        if end_date:
            qs = qs.filter(created_at__date__lte=end_date)
        
        headers = ['NCR Number', 'Product', 'Severity', 'Status', 'Department', 'Process', 'Reported By', 'Created Date', 'Containment Action', 'Root Cause']
        rows = [
            [
                n.ncr_number,
                n.product,
                n.get_severity_display(),
                n.get_status_display(),
                n.department.name if n.department else '',
                n.process.code if n.process else '',
                n.reported_by.full_name if n.reported_by else '',
                n.created_at.strftime('%Y-%m-%d %H:%M'),
                n.containment_action,
                n.root_cause_summary
            ]
            for n in qs
        ]
        return headers, rows, 'NCR_Report'

    elif module == 'capa':
        qs = CAPA.objects.select_related('responsible_person', 'responsible_person__department', 'related_ncr').all()
        if dept_id:
            qs = qs.filter(responsible_person__department_id=dept_id)
        if start_date:
            qs = qs.filter(created_at__date__gte=start_date)
        if end_date:
            qs = qs.filter(created_at__date__lte=end_date)
        
        headers = ['CAPA Number', 'Title', 'Source', 'Status', 'Department', 'Responsible Person', 'Related NCR', 'Due Date', 'Effectiveness Result', 'Created Date']
        rows = [
            [
                c.capa_number,
                c.title,
                c.get_source_display(),
                c.get_status_display(),
                c.responsible_person.department.name if c.responsible_person and c.responsible_person.department else '',
                c.responsible_person.full_name if c.responsible_person else '',
                c.related_ncr.ncr_number if c.related_ncr else '',
                c.due_date.strftime('%Y-%m-%d') if c.due_date else '',
                c.effectiveness_result,
                c.created_at.strftime('%Y-%m-%d')
            ]
            for c in qs
        ]
        return headers, rows, 'CAPA_Report'

    elif module == 'actions':
        qs = Action.objects.select_related('assigned_to', 'department').all()
        if dept_id:
            qs = qs.filter(department_id=dept_id)
        if start_date:
            qs = qs.filter(created_at__date__gte=start_date)
        if end_date:
            qs = qs.filter(created_at__date__lte=end_date)
        
        headers = ['Action Number', 'Title', 'Source Type', 'Source ID', 'Assignee', 'Department', 'Priority', 'Status', 'Due Date', 'Completed Date']
        rows = [
            [
                a.action_number,
                a.title,
                a.get_source_type_display(),
                a.source_id or '',
                a.assigned_to.full_name if a.assigned_to else '',
                a.department.name if a.department else '',
                a.get_priority_display(),
                a.get_status_display(),
                a.due_date.strftime('%Y-%m-%d') if a.due_date else '',
                a.completed_at.strftime('%Y-%m-%d') if a.completed_at else ''
            ]
            for a in qs
        ]
        return headers, rows, 'Actions_Report'

    elif module == 'risks':
        qs = Risk.objects.select_related('department', 'process', 'responsible_person').all()
        if dept_id:
            qs = qs.filter(department_id=dept_id)
        
        headers = ['Risk ID', 'Description', 'Process', 'Likelihood (1-5)', 'Severity (1-5)', 'Score', 'Risk Level', 'Status', 'Department', 'Responsible Person', 'Existing Control']
        rows = [
            [
                r.risk_id,
                r.description,
                r.process.code if r.process else '',
                r.likelihood,
                r.severity,
                r.risk_score,
                r.get_risk_level_display(),
                r.get_status_display(),
                r.department.name if r.department else '',
                r.responsible_person.full_name if r.responsible_person else '',
                r.existing_control
            ]
            for r in qs
        ]
        return headers, rows, 'Risks_Register_Report'

    elif module == 'calibration':
        qs = Equipment.objects.select_related('department').all()
        if dept_id:
            qs = qs.filter(department_id=dept_id)
        
        headers = ['Equipment ID', 'Name', 'Model', 'Serial Number', 'Department', 'Location', 'Interval (Months)', 'Last Calibration', 'Next Calibration', 'Status']
        rows = [
            [
                e.equipment_id,
                e.equipment_name,
                e.model,
                e.serial_number,
                e.department.name if e.department else '',
                e.location,
                e.calibration_frequency_months,
                e.last_calibration_date.strftime('%Y-%m-%d') if e.last_calibration_date else '',
                e.next_calibration_date.strftime('%Y-%m-%d') if e.next_calibration_date else '',
                e.get_status_display()
            ]
            for e in qs
        ]
        return headers, rows, 'Calibration_Equipment_Report'

    elif module == 'training':
        qs = TrainingRecord.objects.select_related('employee', 'course').all()
        if dept_id:
            qs = qs.filter(employee__department_id=dept_id)
        
        headers = ['Employee', 'Employee ID', 'Course Code', 'Course Title', 'Trainer', 'Training Date', 'Expiry Date', 'Result', 'Competency']
        rows = [
            [
                t.employee.full_name,
                t.employee.username,
                t.course.code,
                t.course.title,
                t.trainer,
                t.training_date.strftime('%Y-%m-%d'),
                t.expiry_date.strftime('%Y-%m-%d') if t.expiry_date else 'No Expiry',
                t.get_result_display(),
                t.get_competency_status_display()
            ]
            for t in qs
        ]
        return headers, rows, 'Training_Competency_Report'

    elif module == 'inspections':
        qs = Inspection.objects.select_related('inspector', 'process').all()
        
        headers = ['Inspection Number', 'Lot / Batch Number', 'Type', 'Inspector', 'Inspection Date', 'Result', 'Total Inspected', 'Defective Qty']
        rows = [
            [
                i.inspection_number,
                i.lot_number,
                i.get_inspection_type_display(),
                i.inspector.full_name if i.inspector else '',
                i.inspection_date.strftime('%Y-%m-%d'),
                i.get_result_display(),
                i.sample_size,
                i.defective_quantity
            ]
            for i in qs
        ]
        return headers, rows, 'Inspections_Report'

    else:
        return ['Info'], [['No data']], 'Report'
